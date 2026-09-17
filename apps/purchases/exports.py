"""PDF and Excel exports for purchase documents (wet-ink filing)."""
from io import BytesIO

from django.http import HttpResponse
from openpyxl import Workbook
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.units import mm
from reportlab.platypus import (
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)
from reportlab.lib.styles import getSampleStyleSheet

from apps.purchases.models import ProcessEvent
from apps.purchases.services import list_pr_trail


def _trail_for(obj):
    pr = getattr(obj, "purchase_request", None)
    if pr is None and hasattr(obj, "purchase_order"):
        pr = obj.purchase_order.purchase_request
    if obj.__class__.__name__ == "PurchaseRequest":
        pr = obj
    if pr is None:
        return []
    return list(list_pr_trail(pr=pr))


def _latest_hash(events):
    if not events:
        return ""
    return events[-1].content_hash


def _signature_boxes(kind: str) -> list[str]:
    if kind == "purchase_request":
        return ["Requested by", "Approved by"]
    if kind == "purchase_order":
        return ["Prepared by (operations)", "Vendor acknowledgement (offline)"]
    if kind == "invoice":
        return ["Forwarded for payment"]
    if kind == "warehouse_receipt":
        return ["Received by (warehouse)"]
    return ["Prepared by (operations)"]


def _lines_rows(kind, obj):
    if kind == "purchase_request":
        return [["Description", "Qty", "Unit"]] + [
            [line.description, str(line.quantity), line.unit] for line in obj.lines.all()
        ]
    if kind == "rfq":
        header = ["Line", "Vendor", "Unit price"]
        rows = [header]
        for line in obj.purchase_request.lines.all():
            quotes = line.quotes.select_related("quote__quote_request_vendor__vendor")
            for q in quotes:
                mark = ""
                if line.awarded_vendor_id == q.quote.quote_request_vendor.vendor_id:
                    mark = " (selected)"
                rows.append(
                    [
                        line.description,
                        q.quote.quote_request_vendor.vendor.slug + mark,
                        str(q.unit_price),
                    ]
                )
        return rows
    if kind in ("purchase_order", "invoice"):
        po = obj if kind == "purchase_order" else obj.purchase_order
        return [["Description", "Qty", "Unit", "Unit price"]] + [
            [line.description, str(line.quantity), line.unit, str(line.unit_price)]
            for line in po.lines.all()
        ]
    if kind == "warehouse_receipt":
        return [["Description", "Qty", "Unit"]] + [
            [line.description, str(line.quantity), line.unit] for line in obj.lines.all()
        ]
    return [["—"]]


def _title(kind, obj) -> str:
    mapping = {
        "purchase_request": f"Purchase request {obj.slug}",
        "rfq": f"Quote comparison {obj.slug}",
        "purchase_order": f"Purchase order {obj.slug}",
        "invoice": f"Invoice {obj.slug}",
        "warehouse_receipt": f"Warehouse receipt {obj.slug}",
    }
    return mapping.get(kind, obj.slug)


def build_pdf(kind, obj) -> bytes:
    """Render a PDF with line table, wet-ink boxes, and signed trail."""
    buf = BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, title=_title(kind, obj))
    styles = getSampleStyleSheet()
    story = []
    org = obj.org if hasattr(obj, "org") else getattr(obj, "purchase_order").org
    story.append(Paragraph(org.name, styles["Title"]))
    story.append(Paragraph(_title(kind, obj), styles["Heading2"]))
    space = getattr(obj, "space", None)
    if space is None and hasattr(obj, "purchase_request"):
        space = obj.purchase_request.space
    if space is not None:
        story.append(Paragraph(f"Space: {space.name}", styles["Normal"]))
    story.append(Spacer(1, 6 * mm))
    data = _lines_rows(kind, obj)
    table = Table(data, hAlign="LEFT")
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#eee")),
                ("GRID", (0, 0), (-1, -1), 0.4, colors.grey),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("LEFTPADDING", (0, 0), (-1, -1), 6),
                ("RIGHTPADDING", (0, 0), (-1, -1), 6),
            ]
        )
    )
    story.append(table)
    story.append(Spacer(1, 10 * mm))
    story.append(Paragraph("Manual signature (offline filing)", styles["Heading3"]))
    for label in _signature_boxes(kind):
        box = Table(
            [
                [Paragraph(f"<b>{label}</b>", styles["Normal"])],
                ["Signature: ________________________________"],
                ["Name: ______________________________________"],
                ["Designation / role: ________________________"],
                ["Date: ______________________________________"],
                [Spacer(1, 18 * mm)],
            ],
            colWidths=[170 * mm],
        )
        box.setStyle(
            TableStyle(
                [
                    ("BOX", (0, 0), (-1, -1), 1, colors.black),
                    ("LEFTPADDING", (0, 0), (-1, -1), 8),
                    ("TOPPADDING", (0, 0), (-1, -1), 4),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
                ]
            )
        )
        story.append(box)
        story.append(Spacer(1, 6 * mm))
    events = _trail_for(obj)
    story.append(Paragraph("Signed digital trail", styles["Heading3"]))
    trail_rows = [["Action", "Actor", "When", "Hash"]]
    for event in events:
        trail_rows.append(
            [
                event.action,
                event.actor_slug,
                event.created_at.isoformat(),
                event.content_hash[:16],
            ]
        )
    if len(trail_rows) == 1:
        trail_rows.append(["—", "—", "—", "—"])
    trail = Table(trail_rows, hAlign="LEFT")
    trail.setStyle(
        TableStyle(
            [
                ("GRID", (0, 0), (-1, -1), 0.3, colors.grey),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("FONTSIZE", (0, 0), (-1, -1), 8),
            ]
        )
    )
    story.append(trail)
    latest = _latest_hash(events)
    story.append(Spacer(1, 4 * mm))
    story.append(Paragraph(f"Document hash: {latest}", styles["Normal"]))
    doc.build(story)
    return buf.getvalue()


def build_xlsx(kind, obj) -> bytes:
    """Workbook with Lines and Trail sheets."""
    wb = Workbook()
    lines_ws = wb.active
    lines_ws.title = "Lines"
    for row in _lines_rows(kind, obj):
        lines_ws.append(row)
    trail_ws = wb.create_sheet("Trail")
    trail_ws.append(["action", "actor_slug", "created_at", "content_hash", "signature"])
    events = _trail_for(obj)
    for event in events:
        trail_ws.append(
            [
                event.action,
                event.actor_slug,
                event.created_at.isoformat(),
                event.content_hash,
                event.signature,
            ]
        )
    meta = wb.create_sheet("Meta")
    meta.append(["latest_content_hash", _latest_hash(events)])
    meta.append(["document", _title(kind, obj)])
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def export_document(*, kind: str, obj, fmt: str) -> HttpResponse:
    """Return an attachment HTTP response for pdf or xlsx."""
    if fmt == "pdf":
        payload = build_pdf(kind, obj)
        content_type = "application/pdf"
        filename = f"{obj.slug}.pdf"
    else:
        payload = build_xlsx(kind, obj)
        content_type = (
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
        filename = f"{obj.slug}.xlsx"
    response = HttpResponse(payload, content_type=content_type)
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    return response
